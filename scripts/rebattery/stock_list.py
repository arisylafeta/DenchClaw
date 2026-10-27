#!/usr/bin/env python3
"""Build a buyer-facing ReBattery stock list (.xlsx) from DenchClaw stock and bulk lots.

  stock_list.py --out FILE.xlsx [--chemistry LFP,NMC] [--type pack,module,cell,container]
                [--make Volkswagen,Tesla] [--min-kwh N] [--max-kwh N]
                [--countries "United Kingdom,Germany"] [--min-qty N]
                [--source marketplace,dismantler,bulk] [--prices] [--currency GBP|EUR]
                [--title "ReBattery LFP stock"] [--limit N] [--dry-run]

Sources (local Postgres `denchclaw`, read-only transaction):
  crm_stock_opportunities  item-level stock. "marketplace" = rows linked to a published marketplace
                           listing; "dismantler" = every other available row.
  crm_bulk_listings        kind='supply', status open/qualifying ("bulk").

Output contract (Alex, 27 Oct 2026): one row per lot on sheet "Stock", fixed columns, country only,
no supplier names, towns, supplier refs, VINs, links, contacts, costs or per-kWh prices. Prices only
with --prices: per unit, marketplace listings only. A leak check scans every output cell against the
private values of the source rows and fails hard (exit 2, no file) on any hit.

Prints a JSON summary to stdout. --dry-run prints the summary without writing the file.
"""

from __future__ import annotations

import argparse
import collections
import datetime as dt
import hashlib
import json
import re
import sys

STOCK_COLUMNS = [
    "Ref", "Type", "Make", "Model / product", "Chemistry", "Format", "kWh per unit", "Ah", "V",
    "Qty", "Total kWh", "SoH %", "Condition / test status", "Year", "Country", "Availability",
]
TYPES = ["Pack", "Module", "Cell", "Container"]
SOURCES = ["marketplace", "dismantler", "bulk"]
NOTE_TEXT = ("Figures are supplier-stated unless tested; subject to availability and inspection. "
             "Location shown by country only.")

# --- capacity classes, copied from skills/rebattery-commercial/buyer-inventory-csv -----------------
# Stated kWh -> (gross, usable). Only classes where the stated figure is unambiguous.
CAPACITY = {
    "MEB": {55: (55, 52), 58: (62, 58), 62: (62, 58), 63: (63, 59),
            77: (82, 77), 82: (82, 77), 83: (82, 77)},
    "MLB evo": {64: (71, 64), 71: (71, 64), 86: (95, 86), 95: (95, 86),
                106: (114, 106), 114: (114, 106)},
}
# Gross classes that are single-chemistry, so a blank chemistry can be filled.
NMC_ONLY = {"MEB": {62, 63, 82, 84}, "MLB evo": {71, 95, 114}}
VW_GROUP = {"Volkswagen", "Skoda", "Cupra", "SEAT", "Audi"}

# --- makes -------------------------------------------------------------------------------------------
MAKES = {
    "VOLKSWAGEN": "Volkswagen", "VW": "Volkswagen", "SKODA": "Skoda", "ŠKODA": "Skoda",
    "CUPRA": "Cupra", "SEAT": "SEAT", "AUDI": "Audi", "TESLA": "Tesla", "BYD": "BYD", "MG": "MG",
    "BMW": "BMW", "MINI": "MINI", "KIA": "Kia", "HYUNDAI": "Hyundai", "NISSAN": "Nissan",
    "RENAULT": "Renault", "PEUGEOT": "Peugeot", "CITROEN": "Citroen", "CITROËN": "Citroen",
    "OPEL": "Opel", "VAUXHALL": "Vauxhall", "MERCEDES": "Mercedes-Benz",
    "MERCEDES-BENZ": "Mercedes-Benz", "POLESTAR": "Polestar", "VOLVO": "Volvo", "TOYOTA": "Toyota",
    "LEXUS": "Lexus", "HONDA": "Honda", "FORD": "Ford", "JAGUAR": "Jaguar", "LAND ROVER": "Land Rover",
    "PORSCHE": "Porsche", "FIAT": "Fiat", "JEEP": "Jeep", "DS": "DS", "MAXUS": "Maxus",
    "MITSUBISHI": "Mitsubishi", "MAZDA": "Mazda", "SMART": "smart", "DACIA": "Dacia",
    "SUBARU": "Subaru", "SUZUKI": "Suzuki", "BENTLEY": "Bentley", "IVECO": "Iveco", "FPT": "FPT",
    "STELLANTIS": "Stellantis", "PROTERRA": "Proterra", "REPT": "REPT", "LG": "LG",
    "LG ENERGY SOLUTION": "LG Energy Solution", "LG CHEM": "LG Chem", "PANASONIC": "Panasonic",
    "BAK": "BAK", "KOKAM": "Kokam", "CATL": "CATL", "SAMSUNG SDI": "Samsung SDI", "SK ON": "SK On",
    "EVE": "EVE", "CALB": "CALB", "GOTION": "Gotion", "EP ENERGY": "EP Energy",
    "RELYION": "Relyion", "NEC": "NEC", "SUNGROW": "Sungrow", "HUAWEI": "Huawei", "LEVC": "LEVC",
    "ZERO": "Zero", "ENVISION AESC": "Envision AESC", "AESC": "AESC", "TIER ONE": "",
    "VOLKWAGEN": "Volkswagen", "LYNK": "Lynk & Co", "LYNK & CO": "Lynk & Co", "XPENG": "XPeng",
    "SAMSUNG": "Samsung SDI", "PYLONTECH": "Pylontech", "PYLON TECHNOLOGIES CO., LTD.": "Pylontech",
    "PYLON TECHNOLOGIES": "Pylontech", "MG4": "MG", "DS AUTOMOBILES": "DS", "NIO": "NIO",
    "CHEVROLET": "Chevrolet", "HONGQI": "Hongqi", "SERES": "Seres", "VOYAH": "Voyah",
}
MAKE_ALIASES = {"vw": "volkswagen", "mercedes": "mercedes-benz", "citroën": "citroen", "škoda": "skoda"}

# --- countries ---------------------------------------------------------------------------------------
COUNTRY_NAMES = [
    "United Kingdom", "United States", "Netherlands", "Germany", "Belgium", "France", "Spain", "Italy",
    "Poland", "Sweden", "Norway", "Denmark", "Finland", "Lithuania", "Latvia", "Estonia", "Ireland",
    "Portugal", "Austria", "Switzerland", "Czechia", "Slovakia", "Hungary", "Romania", "Bulgaria",
    "Greece", "Turkey", "Ukraine", "Slovenia", "Croatia", "Serbia", "Luxembourg", "China", "Canada",
    "Mexico", "Japan", "South Korea", "India", "Australia", "United Arab Emirates", "Israel", "Iceland",
]
COUNTRY_ALIASES = {
    "uk": "United Kingdom", "great britain": "United Kingdom", "england": "United Kingdom",
    "scotland": "United Kingdom", "wales": "United Kingdom", "northern ireland": "United Kingdom",
    "england and wales": "United Kingdom", "usa": "United States", "united states of america": "United States",
    "czech republic": "Czechia", "holland": "Netherlands", "the netherlands": "Netherlands",
    "deutschland": "Germany", "türkiye": "Turkey", "uae": "United Arab Emirates",
}
# ISO codes only match as an exact upper-case token.
COUNTRY_CODES = {
    "GB": "United Kingdom", "UK": "United Kingdom", "US": "United States", "USA": "United States",
    "NL": "Netherlands", "DE": "Germany", "BE": "Belgium", "FR": "France", "ES": "Spain", "IT": "Italy",
    "PL": "Poland", "SE": "Sweden", "NO": "Norway", "DK": "Denmark", "FI": "Finland", "LT": "Lithuania",
    "LV": "Latvia", "EE": "Estonia", "IE": "Ireland", "PT": "Portugal", "AT": "Austria",
    "CH": "Switzerland", "CZ": "Czechia", "SK": "Slovakia", "HU": "Hungary", "RO": "Romania",
    "TR": "Turkey", "UA": "Ukraine", "CN": "China",
}
_COUNTRY_WORDS = {c.lower(): c for c in COUNTRY_NAMES}
_COUNTRY_WORDS.update(COUNTRY_ALIASES)
_COUNTRY_RE = re.compile(r"(?<![\w])(" + "|".join(
    re.escape(k) for k in sorted(_COUNTRY_WORDS, key=len, reverse=True)) + r")(?![\w])", re.I)
_CODE_RE = re.compile(r"(?<![\w])(" + "|".join(COUNTRY_CODES) + r")(?![\w])")

# Words that appear in company names but say nothing about who they are.
GENERIC_WORDS = {
    "limited", "ltd", "auto", "autos", "autohandel", "recycling", "recyclers", "parts", "car", "cars",
    "supply", "supplies", "group", "the", "and", "nv", "bv", "gmbh", "sl", "sa", "uk", "electric",
    "ev", "battery", "batteries", "energy", "power", "solutions", "systems", "storage", "motors",
    "motor", "trading", "services", "service", "global", "europe", "european", "international",
    "technologies", "technology", "green", "metals", "salvage", "breakers", "spares", "dismantlers",
    "dismantling", "demontage", "bildemontering", "bildelar", "onderdelen", "automotive", "company",
    "industries", "industrial", "holding", "holdings", "advisory", "seller", "jules", "north", "south",
    "east", "west", "branch", "street", "road", "house", "park", "unit", "site", "warehouse",
    "district", "county", "region", "city", "centre", "center", "limited.", "co.", "inc", "llc",
    "corp", "corporation", "partners", "kft", "kft.", "ab", "as", "oy", "aps", "srl", "spa", "sp",
    "z.o.o", "ltd.", "packs", "pack", "module", "modules", "cells", "cell", "used", "new", "tier",
    "one", "stock", "lithium", "hybrid", "vehicle", "vehicles", "mobility", "circular", "second",
    "life", "renewable", "renewables", "store", "shop", "garage", "sales", "export", "import",
    "logistics", "traders", "trade", "direct", "online", "ebay", "station", "e-mobility", "supplier",
    "suppliers", "not stated", "stated", "unknown", "given", "country", "europe (not stated)",
}

FORMAT_WORDS = [("Cylindrical", r"cylindrical|\b18650\b|\b21700\b|\b2170\b|\b4680\b|n18650|ncr2170|ncr18650"),
                ("Pouch", r"\bpouch\b"), ("Prismatic", r"\bprismatic\b")]

PRIVATE_DETAIL_REF_KEYS = (
    "stock_ref", "supplier_item_number", "dalys_code", "internet_id", "bildelsbasen_listing_id",
    "site_part_id", "donor_stock_no", "tag_no", "donor_vehicle_no", "supplier_vehicle_no",
    "onderdelenlijn_onderdeelid", "onderdelenlijn_vehicle_no", "supplier_ref_onderdelenlijn",
    "autoplius_id", "rrr_sku", "onderdelenlijn_sku", "japoto_internal_sku", "previous_supplier_ref",
    "quote_no", "vin", "vin_masked", "serial", "gbd_part_no", "supplier_stock_id_candidate",
)
PRIVATE_DETAIL_TOWN_KEYS = ("branch", "gbd_branch", "warehouse_location", "storage_location",
                            "location_raw")
PRIVATE_DETAIL_NAME_KEYS = ("supplier_as_listed", "seller_contact")


# =================================================================================================
# Normalisation helpers
# =================================================================================================

def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        s = str(v).strip().replace(",", "").rstrip("%").strip()
        return float(s) if s else None
    except (TypeError, ValueError):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None


def norm_chemistry(text) -> str:
    """Normalise a free-text chemistry to LFP / NMC [nnn] / NCA / LMO / LTO / NiMH / Li-ion.

    Several families named at once ("LFP | NCA | NMC") is ambiguous and returns blank.
    """
    if not text:
        return ""
    t = str(text).strip()
    low = t.lower()
    fams = []
    if re.search(r"\blfp\b|lifepo4|lithium iron phosphate|\blifep", low):
        fams.append("LFP")
    nmc = re.search(r"\bnmc\s?[-/]?\s?(\d{3})\b|\bnmc(\d{3})\b", low)
    if re.search(r"\bnmc|\bncm|ni-mn-co|nickel manganese cobalt", low):
        fams.append("NMC")
    if re.search(r"\bnca\b|nickel cobalt alumin", low):
        fams.append("NCA")
    if re.search(r"\blmo\b|manganese-based|lithium manganese oxide", low):
        fams.append("LMO")
    if re.search(r"\blto\b|titanate", low):
        fams.append("LTO")
    if re.search(r"\bnimh\b|ni-mh|nickel.metal", low):
        fams.append("NiMH")
    if len(fams) > 1:
        # "LMO | NMC" and "LFP | NCA | NMC" style candidate lists: not settled.
        return ""
    if fams == ["NMC"] and nmc:
        return "NMC " + (nmc.group(1) or nmc.group(2))
    if fams:
        return fams[0]
    if re.search(r"li-?ion|lithium|li-po|generic_li_ion|polymer", low):
        return "Li-ion"
    return ""


def chem_family(chem: str) -> str:
    return chem.split(" ")[0].upper() if chem else ""


def chem_filter_value(text: str) -> str:
    """User filter term -> family ("LiFePO4" -> "LFP", "NCM" -> "NMC")."""
    c = norm_chemistry(text)
    return chem_family(c) if c else text.strip().upper()


def norm_make(text) -> str:
    if not text:
        return ""
    t = re.sub(r"\s+", " ", str(text)).strip()
    up = t.upper()
    if up in MAKES:
        return MAKES[up]
    if up.split(" ")[0] in MAKES and MAKES[up.split(" ")[0]] in ("Volkswagen", "Mercedes-Benz"):
        return MAKES[up.split(" ")[0]]
    # Legal entity names used as a make ("WeVolt Renewable Energy Technology Co Limited") -> brand word.
    if re.search(r"(?i)\b(co\.?|ltd\.?|limited|inc\.?|gmbh|llc)\s*\.?$", t) and len(t.split()) > 2:
        t = t.split()[0].rstrip(",")
    return t.title() if t.isupper() and len(t) > 3 else t


def make_key(text: str) -> str:
    k = norm_make(text).lower()
    return MAKE_ALIASES.get(k, k)


def find_make(text: str) -> tuple[str, str]:
    """First known make in a title. Returns (make, the matched text)."""
    best = None
    for key in sorted(MAKES, key=len, reverse=True):
        m = re.search(r"(?<![\w])" + re.escape(key) + r"(?![\w])", text, re.I)
        if m and (best is None or m.start() < best[0].start()):
            best = (m, key)
    if not best:
        return "", ""
    return MAKES[best[1]], best[0].group(0)


def countries_in(text) -> list[str]:
    """Country names found in a location string, in order, de-duplicated."""
    if not text:
        return []
    s = str(text)
    found = []
    for m in _COUNTRY_RE.finditer(s):
        found.append((m.start(), _COUNTRY_WORDS[m.group(1).lower()]))
    for m in _CODE_RE.finditer(s):
        found.append((m.start(), COUNTRY_CODES[m.group(1)]))
    out = []
    for _, c in sorted(found):
        if c not in out:
            out.append(c)
    return out


def norm_country(text) -> str:
    return ", ".join(countries_in(text))


def country_filter_value(text: str) -> str:
    c = countries_in(text)
    return c[0] if c else text.strip()


def town_terms(location) -> set[str]:
    """Every non-country part of a location string (towns, sites, streets, postcodes)."""
    terms = set()
    if not location:
        return terms
    s = str(location)
    parts = re.split(r"[,;/]|\band\b|\(|\)|:", s)
    for p in parts:
        p = re.sub(r"\s+", " ", p).strip(" .-")
        if not p or len(p) < 4:
            continue
        if not _COUNTRY_RE.sub("", p).strip() or not _CODE_RE.sub("", p).strip(" ."):
            continue
        pl = p.lower()
        if pl in _COUNTRY_WORDS or pl in GENERIC_WORDS:
            continue
        # strip a leading "second site" style qualifier and country words inside the part
        p2 = re.sub(r"(?i)^(second site|first site|site|exact site|exact collection address)\s+", "", p)
        p2 = _COUNTRY_RE.sub("", p2).strip(" .-")
        if len(p2) >= 4 and p2.lower() not in GENERIC_WORDS:
            terms.add(p2.lower())
            # distinctive single words too: "Kaunas district" -> "kaunas", "Bil1Din North branch" -> "bil1din"
            for w in re.findall(r"[^\W_][\w\-']*", p2.lower()):
                if len(w) >= 4 and not w.isdigit() and w not in GENERIC_WORDS and w.upper() not in MAKES \
                        and w not in _COUNTRY_WORDS:
                    terms.add(w)
    return terms


def name_terms(name) -> set[str]:
    """A supplier/person name, plus its distinctive words (>= 4 chars, not generic, not a make)."""
    terms = set()
    if not name:
        return terms
    n = re.sub(r"\s+", " ", str(name)).strip()
    if len(n) >= 4:
        terms.add(n.lower())
    for raw in re.findall(r"[^\W_][\w\-']*", n):
        w = raw.lower()
        acronym = len(raw) == 3 and raw.isupper() and raw.isalnum()     # GTV, ASM, LKQ
        if (len(w) < 4 and not acronym) or w in GENERIC_WORDS or w.upper() in MAKES or w in _COUNTRY_WORDS:
            continue
        if w.isdigit():
            continue
        terms.add(w)
    return terms


def ref_terms(ref) -> set[str]:
    """A supplier reference and its code-like parts (>= 4 chars with a digit; numbers >= 5)."""
    terms = set()
    if ref is None:
        return terms
    r = str(ref).strip()
    if not r:
        return terms
    if len(r) >= 4:
        terms.add(r.lower())
    for p in re.split(r"[:\s/|]+", r):
        p = p.strip().lower()
        for q in {p, *p.split("-")} if p.count("-") == 1 else {p}:
            if len(q) >= 4 and re.search(r"\d", q) and not (q.isdigit() and len(q) < 5):
                terms.add(q)
    return terms


def year_of(v) -> int | None:
    if v is None:
        return None
    if isinstance(v, (int, float)) and 1990 <= v <= 2035:
        return int(v)
    m = re.fullmatch(r"\s*((?:19|20)\d\d)(?:\.0)?\s*", str(v))
    return int(m.group(1)) if m else None


def soh_of(v) -> float | None:
    n = num(v)
    if n is None:
        return None
    if 0 < n <= 1:
        n *= 100
    return round(n) if 0 < n <= 100 else None


def format_of(*texts) -> str:
    t = " ".join(str(x) for x in texts if x).lower()
    for label, rx in FORMAT_WORDS:
        if re.search(rx, t):
            return label
    return ""


# --- VW Group model parsing, copied from the buyer-inventory-csv skill ------------------------------

def platform_of(model, title, pn):
    text = f"{model} {title}"
    if pn.startswith("0Z1") or re.search(r"\bID\.?\s?\d|enyaq|elroq|born|Q4", text, re.I):
        return "MEB"
    if pn.startswith("4KE") or (re.search(r"e-tron", text, re.I) and not re.search(r"\bGT\b", text)):
        return "MLB evo"
    return ""


def model_variant(brand, model, title, platform):
    """Pack-relevant model and variant. Vehicle trim names are dropped."""
    m = model or title
    if brand == "Volkswagen" and platform == "MEB":
        if re.fullmatch(r"(?i)1st", m.strip()):
            return "ID. series", "1st Edition"
        if re.fullmatch(r"(?i)pro", m.strip()):
            return "ID. series", "Pro"
        mm = re.search(r"ID\.?\s?(\d)", m, re.I)
        variant = "Pro" if re.search(r"\bPro\b", m) else ""
        return (f"ID.{mm.group(1)}" if mm else ""), variant
    if brand == "Skoda" and platform == "MEB":
        if re.search("elroq", m, re.I):
            return "Elroq", ""
        v = re.search(r"(iV \d+x?|VRS)", m, re.I)
        return "Enyaq", (v.group(1).replace("VRS", "vRS").replace("vrs", "vRS") if v else "")
    if brand == "Cupra" and platform == "MEB":
        return "Born", ""
    if brand == "Audi" and platform == "MEB":
        v = re.search(r"\b(35|40|45|50)\b", m)
        q = " quattro" if v and re.search("quattro", m, re.I) else ""
        return "Q4 e-tron", (v.group(1) + q if v else "")
    if brand == "Audi" and platform == "MLB evo":
        if re.match(r"(?i)Q8", m):
            return "Q8 e-tron", "55 quattro"
        base = "e-tron Sportback" if "sportback" in m.lower() else (
            "e-tron S" if re.search(r"e-tron S\b", m, re.I) else "e-tron")
        v = re.search(r"\b(50|55)\b", m)
        return base, (v.group(1) + " quattro" if v else ("quattro" if base == "e-tron S" else ""))
    return None, None


KEEP_UPPER = {"PHEV", "FHEV", "MHEV", "HEV", "BEV", "AWD", "RWD", "FWD", "GTI", "GTE", "SUV", "BESS",
              "CATL", "NMC", "LFP", "EV"}


def clean_model(make: str, model: str) -> str:
    """Vehicle model / product without make, years, kWh, part numbers in brackets or pack words."""
    if not model:
        return ""
    m = str(model)
    m = re.sub(r"\([^)]*\)", " ", m)                     # bracketed part numbers, notes
    m = re.sub(r"(?i)\b(19|20)\d\d\s*-\s*(?:\d{2,4}\b)?|\b\d{2}\s*-?>+|\b(19|20)\d\d\b", " ", m)
    m = re.sub(r"(?i)\b\d+(?:[.,]\d+)?\s*-\s*\d+(?:[.,]\d+)?\s*[km]wh\b", " ", m)
    m = re.sub(r"(?i)\b\d+(?:[.,]\d+)?\s*[km]wh\b", " ", m)
    m = re.sub(r"(?i)\b(hv |high[- ]voltage |traction )?(battery|batteries|batterie|batteri|pack|packs|"
               r"module|modules|cells?|other|type unclear|complete|unused|new)\b", " ", m)
    if make:
        keys = [k for k, v in MAKES.items() if v == make and not re.search(r"\d", k)] + [make]
        for k in sorted(set(keys), key=len, reverse=True):
            m = re.sub(r"(?i)(?<![\w])" + re.escape(k) + r"(?![\w])", " ", m)
    m = re.sub(r"\s+", " ", m).strip(" ,;:-/")
    words = []
    for w in m.split(" "):
        core = w.strip(",;:")
        if re.fullmatch(r"[IVX]+", core):
            pass                                            # roman numeral generations: VII, IV
        elif core.isupper() and core.replace("-", "").isalpha() and len(core) > 3 and core not in KEEP_UPPER:
            w = w.replace(core, core.title())
        elif "/" in w and w.isupper():
            w = "/".join(x.title() if x.isalpha() and len(x) > 3 and x not in KEEP_UPPER else x
                         for x in w.split("/"))
        words.append(w)
    return " ".join(words).strip()


# =================================================================================================
# Condition
# =================================================================================================

def condition_from_text(text: str, tested=None, damaged=None, installed=None) -> str:
    t = (text or "").lower()
    base = ""
    if damaged is True or re.search(r"for parts|not working|defect|damage|faulty|\bcore\b", t):
        base = "Damaged / for parts"
    elif re.search(r"end[ _]of[ _]life", t):
        base = "End of life"
    elif re.search(r"never[- ]used|\bunused\b|never in service|\bnew\b|surplus", t) and not re.search(r"as new|like new", t):
        base = "New (unused)"
    elif tested is False or re.search(r"not tested|untested|ej testad|not tested for operation|niet getest", t):
        base = "Used, untested"
    elif tested is True or re.search(r"tested|getest|testad|testat|testprotokoll", t):
        base = "Used, tested"
    elif re.search(r"\b(excellent|great|functional)\b", t):
        g = re.search(r"\b(excellent|great|functional)\b", t).group(1)
        base = f"Used, supplier grade {g.capitalize()}"
    elif re.search(r"\bused\b|naudota|removed|quality|gebraucht|brukt|begagnad", t):
        base = "Used, test status not stated"
    if installed is True or re.search(r"still fitted|in (the )?vehicle", t):
        base = (base + "; still fitted to vehicle") if base else "Still fitted to vehicle"
    return base


def stock_condition(d: dict) -> str:
    c = d.get("condition")
    text = c if isinstance(c, str) else ""
    if isinstance(c, str) and re.fullmatch(r"[A-Za-z]\d?\*?", c.strip()):
        text = ""                                       # bare supplier grade letters: meaning unknown
    for k in ("grade_meaning", "test_claim", "condition_text"):
        if isinstance(d.get(k), str):
            text += " " + d[k]
    tested = d.get("tested") if isinstance(d.get("tested"), bool) else None
    damaged = d.get("damaged") if isinstance(d.get("damaged"), bool) else None
    installed = d.get("installed") if isinstance(d.get("installed"), bool) else None
    return condition_from_text(text, tested, damaged, installed)


# =================================================================================================
# Type inference
# =================================================================================================

def type_word(text: str) -> str:
    t = (text or "").lower()
    if re.search(r"\bcontainers?\b|\bbess\b|\bcabinets?\b|megapack|powerpack|\bsystems?\b", t):
        return "Container"
    if re.search(r"\bmodules?\b", t):
        return "Module"
    if re.search(r"\bcells?\b", t):
        return "Cell"
    if re.search(r"\bpacks?\b", t):
        return "Pack"
    return ""


def explicit_type(v) -> str:
    t = str(v or "").strip().lower()
    return {"pack": "Pack", "module": "Module", "cell": "Cell", "modules": "Module", "cells": "Cell",
            "container": "Container", "battery pack": "Pack", "complete_pack_candidate": "Pack"}.get(t, "")


def stock_type(row: dict) -> str:
    """Pack/Module/Cell, '' when the row says its type is unclear."""
    d = row.get("details") or {}
    mp = d.get("marketplace") or {}
    for v in (mp.get("format"), d.get("format"), row.get("quantity_unit"), d.get("scope"), d.get("kind")):
        t = explicit_type(v)
        if t:
            return t
    title = row.get("title") or ""
    if str(d.get("format") or "").lower() == "other":
        return type_word(title) if type_word(title) in ("Pack", "Module", "Cell") else ""
    w = type_word(title)
    return w if w in ("Module", "Cell") else "Pack"


def bulk_type(row: dict) -> str:
    spec = row.get("spec") or {}
    ctx = row.get("context") or {}
    t = explicit_type(spec.get("format"))
    if t:
        return t
    for text in (ctx.get("listing_title"), row.get("summary")):
        w = type_word(text or "")
        if w:
            return w
    return {"cells": "Cell", "packs": "Pack", "modules": "Module", "systems": "Container"}.get(
        str(spec.get("trade_kind") or "").lower(), "")


# =================================================================================================
# Item builders. Each returns (item | None, exclusion reason | None, private terms).
# =================================================================================================

def _private_from_stock(row: dict) -> dict:
    d = row.get("details") or {}
    names, towns, refs = set(), set(), set()
    for n in (row.get("company_name"), row.get("person_name"), (d.get("ebay") or {}).get("seller"),
              (d.get("evidence") or {}).get("seller") if isinstance(d.get("evidence"), dict) else None,
              *(d.get(k) for k in PRIVATE_DETAIL_NAME_KEYS)):
        if isinstance(n, str):
            names |= name_terms(n)
    for loc in (row.get("location"), row.get("company_city"), *(d.get(k) for k in PRIVATE_DETAIL_TOWN_KEYS)):
        if isinstance(loc, str):
            towns |= town_terms(loc)
    eloc = (d.get("ebay") or {}).get("location") or {}
    for k in ("city", "stateOrProvince", "postalCode"):
        if isinstance(eloc.get(k), str):
            towns |= town_terms(eloc[k])
    refs |= ref_terms(row.get("supplier_ref"))
    for k in PRIVATE_DETAIL_REF_KEYS:
        if d.get(k) not in (None, "", True, False):
            refs |= ref_terms(d.get(k))
    if (d.get("ebay") or {}).get("item_id"):
        refs |= ref_terms(d["ebay"]["item_id"])
    attrs = d.get("attributes") if isinstance(d.get("attributes"), dict) else {}
    for k in ("vin", "ref_nr", "dem_nr", "lagerplass", "vis_orgnr"):
        if attrs.get(k):
            refs |= ref_terms(attrs[k])
    # A ref part that is also the row's own product name ("INR18650-35E", "R135") is a product
    # code, not a supplier code. Pure numbers always stay private.
    mp = d.get("marketplace") if isinstance(d.get("marketplace"), dict) else {}
    own = _tokens(" ".join(str(x) for x in (row.get("title"), d.get("model"), mp.get("model")) if x))
    own |= {q for t in own for q in re.split(r"[\-'.]", t)}
    refs = {r for r in refs if r.isdigit() or r not in own}
    return {"names": names, "towns": towns, "refs": refs}


def withhold_make(make: str, names: set[str]) -> tuple[str, bool]:
    """Blank the make when it is (part of) the row's own supplier name."""
    if not make:
        return make, False
    words = {w for w in re.findall(r"[^\W_][\w\-']*", make.lower()) if len(w) >= 4}
    if make.lower() in names or words & names:
        return "", True
    return make, False


def stock_exclusion(row: dict) -> str | None:
    d = row.get("details") or {}
    if row.get("status") not in ("opportunity", "listed"):
        return "stale" if row.get("status") == "stale" else f"status_{row.get('status')}"
    if row.get("marketplace_status") in ("withdrawn", "completed", "missing"):
        return f"marketplace_{row['marketplace_status']}"
    if d.get("not_a_battery") or str(d.get("scope") or "").lower().startswith("not a battery") \
            or d.get("scope") in ("obvious_nonpack", "spares"):
        return "not_a_traction_battery"
    if d.get("excluded"):
        return "flagged_excluded"
    if d.get("duplicate_of") or d.get("same_pack_as"):
        return "duplicate"
    state = (d.get("availability") or {}).get("state") if isinstance(d.get("availability"), dict) else None
    if state in ("sold", "not_found"):
        return "sold_or_not_found"
    if state == "unclear":
        return "availability_unclear"
    if str(d.get("powertrain") or "").lower() in ("mhev", "mild hybrid") or d.get("scope") == "mild_hybrid_48v":
        return "mild_hybrid_48v"
    text = " ".join(str(x) for x in (row.get("title"), d.get("model"), d.get("powertrain_raw"),
                                      d.get("hybrid_type"), d.get("donor_engine")) if x)
    if re.search(r"(?i)\beTSI\b|mild[- ]?hybrid|\bMHEV\b|48\s?V\s+(mild|hybrid)", text):
        return "mild_hybrid_48v"
    return None


def stock_source(row: dict) -> str:
    return "marketplace" if row.get("marketplace_listing_id") and row.get("marketplace_status") == "published" \
        else "dismantler"


def build_stock_item(row: dict):
    priv = _private_from_stock(row)
    reason = stock_exclusion(row)
    if reason:
        return None, reason, priv
    d = row.get("details") or {}
    mp = d.get("marketplace") if isinstance(d.get("marketplace"), dict) else {}
    typ = stock_type(row)
    if not typ:
        return None, "type_unclear", priv
    make = norm_make(mp.get("manufacturer") or d.get("make") or "")
    title = row.get("title") or ""
    if not make:
        make = find_make(title)[0]
    raw_model = mp.get("model") or d.get("model") or ""
    pn = str(d.get("part_number") or "").upper()
    stated = None
    for k in ("capacity_kwh_gross", "capacity_kwh", "capacity_kwh_stated_pack", "capacity_kwh_listed"):
        stated = num(d.get(k))
        if stated:
            break
    if not stated:
        stated = num(mp.get("pack_kwh"))
    kwh = stated
    chem = norm_chemistry(mp.get("chemistry")) or norm_chemistry(d.get("chemistry"))
    model = ""
    platform = platform_of(raw_model, title, pn) if make in VW_GROUP else ""
    if platform:
        mv = model_variant(make, raw_model, title, platform)
        if mv[0]:
            model = (mv[0] + (" " + mv[1] if mv[1] else "")).strip()
        if typ == "Pack" and stated and round(stated) in CAPACITY.get(platform, {}):
            kwh = float(CAPACITY[platform][round(stated)][0])
        if not chem and kwh and round(kwh) in NMC_ONLY.get(platform, set()):
            chem = "NMC"
    if not model:
        model = clean_model(make, raw_model or title)
    ah = num(d.get("capacity_ah")) or num(d.get("amp_hours")) or num(d.get("amperage_ah"))
    volts = num(d.get("voltage_nominal")) or num(d.get("voltage_v"))
    qty = num(row.get("quantity"))
    if qty is None:
        qty = num(d.get("quantity_available"))
    soh = soh_of(mp.get("soh")) or soh_of(d.get("soh_percent"))
    year = year_of(mp.get("year_manufacture")) or year_of(d.get("year"))
    country = norm_country(row.get("location"))
    if not country:
        eloc = (d.get("ebay") or {}).get("location") or {}
        country = norm_country(eloc.get("country"))
    if not country:
        country = norm_country(row.get("company_country"))
    # Manufacturer selling direct (Pylontech listing Pylontech): the make would name the supplier.
    make, withheld = withhold_make(make, priv["names"])
    fmt = format_of(d.get("cell_spec"), d.get("format_note"), d.get("format_supplier_term"), title)
    slug = d.get("seo_slug") or ""
    item = {
        "_id": str(row["id"]), "_source": stock_source(row), "_slug": slug,
        "_price": num(row.get("price")) if row.get("source") == "marketplace" else None,
        "_currency": (row.get("currency") or "").upper(), "_fx": d.get("fx") if isinstance(d.get("fx"), dict) else None,
        "_make_withheld": withheld,
        "Type": typ, "Make": make, "Model / product": model, "Chemistry": chem, "Format": fmt,
        "kWh per unit": round(kwh, 4) if kwh else None, "Ah": ah, "V": volts,
        "Qty": int(qty) if qty is not None and qty == int(qty) else qty,
        "SoH %": soh, "Condition / test status": stock_condition(d), "Year": year,
        "Country": country, "Availability": "Now",
    }
    return item, None, priv


def _bulk_qty(row: dict, title: str):
    spec = row.get("spec") or {}
    if num(spec.get("quantity")):
        return num(spec.get("quantity"))
    m = re.match(r"\s*(\d[\d,]*)\s*x\s", title or "")
    if m:
        return num(m.group(1))
    q = str(row.get("quantity") or "").strip()
    m = re.match(r"(?:at least\s+|potentially\s+)?(\d[\d,]*)\s*(packs?|modules?|cells?|containers?|"
                 r"systems?|units?|batteries)?\b", q, re.I)
    if m and (m.group(2) or re.fullmatch(r"\d[\d,]*", q)):
        return num(m.group(1))
    return None


def build_bulk_item(row: dict):
    spec = row.get("spec") or {}
    ctx = row.get("context") or {}
    names, towns, refs = set(), set(), set()
    for n in (row.get("party"), row.get("contact"), row.get("company_name"), row.get("person_name"),
              spec.get("seller_name")):
        if isinstance(n, str):
            names |= name_terms(n)
    for loc in (row.get("location"), row.get("company_city")):
        if isinstance(loc, str):
            towns |= town_terms(loc)
    for r in (spec.get("seller_ref"),):
        if r:
            refs |= ref_terms(str(r).split(":", 1)[-1])
    priv = {"names": names, "towns": towns, "refs": refs}

    if row.get("kind") != "supply" or row.get("status") not in ("open", "qualifying"):
        return None, "bulk_closed", priv
    title = ctx.get("listing_title") or ""
    text = " ".join(x for x in (title, row.get("summary") or "", ctx.get("fact_line") or "",
                                spec.get("fact_line") or "") if x)
    if re.search(r"\bfaulty\b|\bdisposal\b|\bwaste\b", title + " " + (row.get("summary") or ""), re.I):
        return None, "bulk_waste_or_disposal", priv
    qty = _bulk_qty(row, title)
    if qty is None:
        return None, "bulk_no_quantity", priv
    typ = bulk_type(row)
    make = norm_make(spec.get("manufacturer") or "")
    matched = ""
    if not make:
        make, matched = find_make(title or row.get("summary") or "")
    make, withheld = withhold_make(make, names)
    if make == "Volkswagen" and re.search(r"\bVW Group\b", title):
        make = "Volkswagen Group"
    # kWh per unit: spec, else from a title that starts "Nx" (then the figure is per unit).
    kwh = num(spec.get("packKwh"))
    if not kwh and re.match(r"\s*\d[\d,]*\s*x\s", title):
        m = re.search(r"(?<![\d.\-])(\d+(?:\.\d+)?)\s?(kWh|MWh)\b", title, re.I)
        if m:
            kwh = float(m.group(1)) * (1000 if m.group(2).lower() == "mwh" else 1)
    model_src = spec.get("model") or re.sub(r"^\s*\d[\d,]*\s*x\s*", "", title) or ""
    model = clean_model(make if make != "Volkswagen Group" else "Volkswagen", model_src)
    model = re.sub(r"(?i)\b\d+(?:\.\d+)?\s?ah\b|\b(NMC|LFP|NCA|LMO|containers?|systems?|units?|"
                   r"prismatic|pouch|cylindrical|group)\b", " ", model)
    model = re.sub(r"\s+", " ", model).strip(" ,;:-/")
    if typ == "Container" and model and "bess" not in model.lower() and "megapack" not in model.lower():
        model += " BESS"
    if typ == "Container" and not model:
        model = "BESS"
    chem = norm_chemistry(spec.get("chemistry")) or norm_chemistry(
        " ".join(sorted(set(re.findall(r"(?i)\b(?:LFP|LiFePO4|NMC\s?\d{3}|NMC|NCM|NCA|LMO|LTO)\b", text)))))
    if spec.get("chemistry") is None and len({chem_family(norm_chemistry(x)) for x in
                                              re.findall(r"(?i)\b(?:LFP|LiFePO4|NMC|NCM|NCA|LMO|LTO)\b", text)}) > 1:
        chem = ""
    ah = None
    m = re.search(r"(?<![\w.])(\d+(?:\.\d+)?)\s?Ah\b", " ".join([title, str(spec.get("model") or ""),
                                                               row.get("summary") or ""]))
    if m:
        ah = float(m.group(1))
    volts = None
    vs = set(re.findall(r"(?<![\w.])(\d+(?:\.\d+)?)\s?V\b", row.get("summary") or ""))
    if len(vs) == 1:
        volts = float(vs.pop())
    soh = soh_of(spec.get("soh"))
    if soh is None:
        m = re.search(r"(?i)\bSoH\s*(?:of\s*)?(\d{2,3})\s*%", text)
        soh = soh_of(m.group(1)) if m else None
    year = year_of(spec.get("year"))
    if year is None:
        m = re.search(r"(?i)\b(?:produced|built|manufactured|unused)(?:\s+in)?\s+((?:19|20)\d\d)\b", text)
        year = int(m.group(1)) if m else None
    sc = str(spec.get("condition") or "").lower()
    if sc in ("new", "unused"):
        cond = "New (unused)"
    elif sc in ("excellent", "great", "functional"):
        cond = f"Used, supplier grade {sc.capitalize()}"
    elif sc in ("end_of_life", "end of life"):
        cond = "End of life"
    else:
        paren = " ".join(re.findall(r"\(([^)]*)\)", title))
        cond = condition_from_text(paren) or condition_from_text(row.get("summary") or "")
        cond = cond.replace("; still fitted to vehicle", "")
    # Bulk lots: the stated lot location only. The supplier's head-office country is not the lot's.
    country = norm_country(row.get("location"))
    avail = "Now"
    hold = row.get("hold_until") or spec.get("hold_until")
    if hold:
        h = hold if isinstance(hold, dt.date) else dt.date.fromisoformat(str(hold)[:10])
        if h > dt.date.today():
            avail = "From " + h.strftime("%b %Y")
    item = {
        "_id": "bulk:" + str(row["id"]), "_source": "bulk", "_slug": "", "_price": None, "_currency": "",
        "_fx": None, "_make_withheld": withheld,
        "Type": typ, "Make": make, "Model / product": model, "Chemistry": chem,
        "Format": format_of(title, spec.get("model"), row.get("summary")),
        "kWh per unit": round(kwh, 4) if kwh else None, "Ah": ah, "V": volts,
        "Qty": int(qty) if qty == int(qty) else qty, "SoH %": soh, "Condition / test status": cond,
        "Year": year, "Country": country, "Availability": avail,
    }
    return item, None, priv


# =================================================================================================
# Filters, refs, sort, totals
# =================================================================================================

def csv_arg(v):
    return [x.strip() for x in v.split(",") if x.strip()] if v else []


def apply_filters(items: list[dict], args) -> tuple[list[dict], collections.Counter]:
    excluded = collections.Counter()
    chem = {chem_filter_value(x) for x in csv_arg(args.chemistry)}
    types = {x.strip().capitalize() for x in csv_arg(args.type)}
    makes = {make_key(x) for x in csv_arg(args.make)}
    countries = {country_filter_value(x).lower() for x in csv_arg(args.countries)}
    sources = {x.lower() for x in csv_arg(args.source)}
    explicit_zero = args.min_kwh is not None and args.min_kwh <= 0
    kept = []
    for it in items:
        k = it.get("kWh per unit")
        if it["Type"] == "Pack" and k is not None and k < 1 and not explicit_zero:
            excluded["pack_under_1kwh"] += 1
            continue
        if sources and it["_source"] not in sources:
            excluded["filtered_source"] += 1
            continue
        if types and it["Type"] not in types:
            excluded["filtered_type"] += 1
            continue
        if chem and chem_family(it["Chemistry"]) not in chem:
            excluded["filtered_chemistry"] += 1
            continue
        if makes and make_key(it["Make"]) not in makes and not (
                "volkswagen" in makes and it["Make"] == "Volkswagen Group"):
            excluded["filtered_make"] += 1
            continue
        if args.min_kwh is not None or args.max_kwh is not None:
            if k is None:
                excluded["filtered_kwh_unknown"] += 1
                continue
            if (args.min_kwh is not None and k < args.min_kwh) or (args.max_kwh is not None and k > args.max_kwh):
                excluded["filtered_kwh"] += 1
                continue
        if countries and not ({c.lower() for c in it["Country"].split(", ") if c} & countries):
            excluded["filtered_country"] += 1
            continue
        if args.min_qty is not None and (it["Qty"] is None or it["Qty"] < args.min_qty):
            excluded["filtered_qty"] += 1
            continue
        kept.append(it)
    return kept, excluded


def assign_refs(items: list[dict]) -> None:
    """RB- + last 4 of the marketplace seo_slug, else a short hash of the source row id. Stable."""
    def base(it, n):
        if it["_slug"] and re.fullmatch(r"[0-9a-z]{4,}", it["_slug"][-n:] or ""):
            return "RB-" + it["_slug"][-n:].upper()
        return "RB-" + hashlib.sha1(it["_id"].encode()).hexdigest()[:n + 2].upper()
    for it in items:
        it["Ref"] = base(it, 4)
    for n in (6, 8, 10):
        counts = collections.Counter(it["Ref"] for it in items)
        dupes = {r for r, c in counts.items() if c > 1}
        if not dupes:
            return
        for it in items:
            if it["Ref"] in dupes:
                it["Ref"] = base(it, n)


def sort_items(items: list[dict]) -> list[dict]:
    order = {t: i for i, t in enumerate(TYPES)}
    return sorted(items, key=lambda x: (order.get(x["Type"], 9), x["Make"] == "", x["Make"].lower(),
                                        x["Model / product"].lower(), -(x["kWh per unit"] or 0),
                                        -(x["SoH %"] or 0), x["Ref"]))


def fill_totals(items: list[dict]) -> None:
    for it in items:
        k, q = it.get("kWh per unit"), it.get("Qty")
        it["Total kWh"] = round(k * q, 3) if k is not None and q is not None else None


def price_for(it: dict, currency: str):
    """Per-unit price in `currency`, marketplace listings only; None when it cannot be shown."""
    p = it.get("_price")
    if it["_source"] != "marketplace" or p is None:
        return None
    cur = it.get("_currency")
    if cur == currency:
        out = p
    elif currency == "EUR" and it.get("_fx") and num(it["_fx"].get("rate")):
        out = p / num(it["_fx"]["rate"])        # fx.rate = row currency per 1 EUR (ECB reference)
    else:
        return None
    return round(min(out, p) if cur == currency else out, 2)


def columns_for(args) -> list[str]:
    return STOCK_COLUMNS + ([f"Price per unit ({args.currency})"] if args.prices else [])


def to_rows(items: list[dict], args) -> list[list]:
    cols = columns_for(args)
    rows = []
    for it in items:
        if args.prices:
            it[cols[-1]] = price_for(it, args.currency)
        rows.append([it.get(c) if it.get(c) not in ("",) else None for c in cols])
    return rows


def totals_row(rows: list[list], cols: list[str]) -> list:
    qi, ki = cols.index("Qty"), cols.index("Total kWh")
    out = [None] * len(cols)
    out[0] = "Total"
    out[qi] = sum(r[qi] for r in rows if isinstance(r[qi], (int, float)))
    out[ki] = round(sum(r[ki] for r in rows if isinstance(r[ki], (int, float))), 3)
    return out


# =================================================================================================
# Leak check
# =================================================================================================

URL_RE = re.compile(r"https?://|www\.|\b[\w-]+\.(?:com|co\.uk|uk|de|nl|se|no|lt|fr|io|net|org|eu|be|pl|dk)\b", re.I)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
PHONE_RE = re.compile(r"\+\d[\d \-().]{7,}\d|\b0\d{2,4}[ \-]\d{3}[ \-]?\d{3,4}\b")
VIN_RE = re.compile(r"\b(?=[A-HJ-NPR-Z0-9]*\d)(?=[A-HJ-NPR-Z0-9]*[A-HJ-NPR-Z])[A-HJ-NPR-Z0-9]{17}\b")
POSTCODE_RE = re.compile(r"\b[A-PR-UWYZ][A-HK-Y]?\d[A-Z\d]? \d[ABD-HJLNP-UW-Z]{2}\b")
PER_KWH_RE = re.compile(r"(?i)(?:£|€|\$|gbp|eur|usd)\s*\d[\d,.]*\s*/\s*kwh|/\s*kwh|per\s+kwh")


def _tokens(s: str) -> set[str]:
    return set(re.findall(r"[^\W_][\w\-'.]*[^\W_]|[^\W_]", s.lower()))


def leak_check(cells: list, private: dict) -> list[str]:
    """Return offenders as "value: reason (term)". Empty list = clean."""
    names = {t for t in private.get("names", set()) if len(t) >= 3}   # 3 only for acronyms (name_terms)
    towns = {t for t in private.get("towns", set()) if len(t) >= 4}
    refs = {t for t in private.get("refs", set()) if len(t) >= 4}
    single = {}
    multi = []
    for kind, terms in (("supplier name", names), ("town/site", towns), ("supplier ref", refs)):
        for t in terms:
            if re.fullmatch(r"[^\W_][\w\-'.]*[^\W_]|[^\W_]", t):
                single.setdefault(t, kind)
            else:
                multi.append((t, kind))
    offenders = []
    for v in sorted({str(c) for c in cells if c not in (None, "")}):
        low = v.lower()
        toks = _tokens(v)
        # also look at hyphen/dot-separated pieces so "foo-bar" catches "foo"
        pieces = set()
        for tk in toks:
            pieces |= set(re.split(r"[\-'.]", tk))
        for tk in sorted((toks | pieces) & set(single)):
            if single[tk] == "supplier ref" and isinstance(v, str) and v.replace(".", "").isdigit():
                continue
            offenders.append(f"{v!r}: {single[tk]} ({tk})")
        for t, kind in multi:
            if t in low and re.search(r"(?<![\w])" + re.escape(t) + r"(?![\w])", low):
                offenders.append(f"{v!r}: {kind} ({t})")
        if re.fullmatch(r"-?\d+(\.\d+)?", v):
            continue                                      # plain numbers: no patterns apply
        for rx, why in ((URL_RE, "url"), (EMAIL_RE, "email"), (PHONE_RE, "phone"), (VIN_RE, "VIN"),
                        (POSTCODE_RE, "postcode"), (PER_KWH_RE, "per-kWh price")):
            if rx.search(v if rx is not VIN_RE else v.upper()):
                offenders.append(f"{v!r}: {why}")
    return offenders


# =================================================================================================
# Workbook
# =================================================================================================

WIDTHS = {"Ref": 10, "Type": 11, "Make": 16, "Model / product": 28, "Chemistry": 11, "Format": 12,
          "kWh per unit": 12, "Ah": 8, "V": 8, "Qty": 10, "Total kWh": 12, "SoH %": 8,
          "Condition / test status": 34, "Year": 7, "Country": 18, "Availability": 13}
FORMATS = {"kWh per unit": "#,##0.0##", "Total kWh": "#,##0.0", "Ah": "#,##0.#", "V": "#,##0.#",
           "Qty": "#,##0", "SoH %": "0", "Year": "0"}


def notes_lines(args, n_rows: int, total_kwh: float, generated: str) -> list[list]:
    filt = []
    for label, val in (("Chemistry", args.chemistry), ("Type", args.type), ("Make", args.make),
                       ("Min kWh per unit", args.min_kwh), ("Max kWh per unit", args.max_kwh),
                       ("Countries", args.countries), ("Min qty", args.min_qty), ("Source", args.source),
                       ("Limit", args.limit)):
        if val not in (None, ""):
            filt.append(f"{label}: {val}")
    lines = [[args.title or "ReBattery stock list"], [], ["Generated", generated],
             ["Filters used", "; ".join(filt) if filt else "None (all available stock)"],
             ["Lots", n_rows], ["Total kWh (where known)", round(total_kwh, 1)]]
    if args.prices:
        lines.append(["Prices", f"Per unit in {args.currency}, current marketplace listings only; "
                                "blank where no listed price applies"])
    lines += [[], [NOTE_TEXT]]
    return lines


def write_workbook(path: str, cols: list[str], rows: list[list], total: list, notes: list[list], title: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(cols)
    for r in rows:
        ws.append(r)
    ws.append(total)
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="E8EEF4")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for c in ws[ws.max_row]:
        c.font = bold
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(cols))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(rows) + 1)}"
    for i, name in enumerate(cols, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = WIDTHS.get(name, 16)
        fmt = FORMATS.get(name) or ("#,##0.00" if name.startswith("Price per unit") else None)
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].number_format = fmt
    ns = wb.create_sheet("Notes")
    for line in notes:
        ns.append(line)
    ns["A1"].font = Font(bold=True, size=12)
    ns.column_dimensions["A"].width = 26
    ns.column_dimensions["B"].width = 90
    wb.properties.title = title or "ReBattery stock list"
    wb.properties.creator = "ReBattery"
    wb.save(path)


# =================================================================================================
# Data loading (read-only)
# =================================================================================================

STOCK_SQL = """
select s.id, s.status, s.title, s.description, s.quantity, s.quantity_unit, s.price, s.currency,
       s.price_unit, s.location, s.supplier_ref, s.marketplace_listing_id, s.marketplace_status,
       s.source, s.details, c.name as company_name, c.country as company_country,
       c.city as company_city, p.full_name as person_name
  from crm_stock_opportunities s
  left join crm_companies c on c.id = s.company_id
  left join crm_people p on p.id = s.person_id
"""
BULK_SQL = """
select b.id, b.kind, b.party, b.contact, b.summary, b.quantity, b.location, b.status, b.spec,
       b.context, b.hold_until, c.name as company_name, c.country as company_country,
       c.city as company_city, p.full_name as person_name
  from crm_bulk_listings b
  left join crm_companies c on c.id = b.company_id
  left join crm_people p on p.id = b.person_id
 where b.kind = 'supply'
"""


def load_rows(dsn: str):
    import psycopg2
    import psycopg2.extras
    conn = psycopg2.connect(dsn)
    try:
        conn.set_session(readonly=True)
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute("set transaction read only")
            cur.execute(STOCK_SQL)
            stock = [dict(r) for r in cur.fetchall()]
            cur.execute(BULK_SQL)
            bulk = [dict(r) for r in cur.fetchall()]
        conn.rollback()
    finally:
        conn.close()
    return stock, bulk


# =================================================================================================
# Main
# =================================================================================================

def build(stock: list[dict], bulk: list[dict], args) -> dict:
    """Pure core: rows in, everything needed to write + summarise out."""
    items, excluded = [], collections.Counter()
    private = {"names": set(), "towns": set(), "refs": set()}
    for row, builder in [(r, build_stock_item) for r in stock] + [(r, build_bulk_item) for r in bulk]:
        item, reason, priv = builder(row)
        for k in private:
            private[k] |= priv[k]
        if reason:
            excluded[reason] += 1
        else:
            items.append(item)
    # never let a make, country or fixed vocabulary word count as a private term
    safe = {v.lower() for v in MAKES.values() if v} | set(_COUNTRY_WORDS) | {
        "pack", "module", "cell", "container", "now", "used", "new", "tested", "untested"}
    for k in private:
        private[k] -= safe
    # Ref parts that are also product names somewhere in the stock ("INR18650-35E" cells listed with
    # the cell model as their ref, "R135" motor variants) are product codes. Pure numbers stay private.
    product = set()
    for r in stock:
        d = r.get("details") or {}
        mp = d.get("marketplace") if isinstance(d.get("marketplace"), dict) else {}
        product |= _tokens(" ".join(str(x) for x in (r.get("title"), d.get("model"), mp.get("model")) if x))
    product |= {q for t in product for q in re.split(r"[\-'.]", t)}
    private["refs"] = {r for r in private["refs"] if r.isdigit() or r not in product}
    kept, filt = apply_filters(items, args)
    excluded.update(filt)
    assign_refs(kept)
    fill_totals(kept)
    kept = sort_items(kept)
    if args.limit is not None and len(kept) > args.limit:
        excluded["over_limit"] += len(kept) - args.limit
        kept = kept[: args.limit]
    cols = columns_for(args)
    rows = to_rows(kept, args)
    total = totals_row(rows, cols)
    gaps = []
    for i, c in enumerate(cols):
        blank = sum(1 for r in rows if r[i] in (None, ""))
        if blank:
            gap = {"column": c, "blank": blank}
            if c.startswith("Price per unit"):
                by = collections.Counter(kept[j]["_source"] for j, r in enumerate(rows) if r[i] is None)
                gap["by_source"] = dict(by)
            gaps.append(gap)
    withheld = sum(1 for it in kept if it.get("_make_withheld"))
    if withheld:
        gaps.append({"column": "Make", "withheld_supplier_is_manufacturer": withheld})
    return {"items": kept, "cols": cols, "rows": rows, "total": total, "excluded": excluded,
            "gaps": gaps, "private": private}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out")
    ap.add_argument("--chemistry")
    ap.add_argument("--type")
    ap.add_argument("--make")
    ap.add_argument("--min-kwh", type=float)
    ap.add_argument("--max-kwh", type=float)
    ap.add_argument("--countries")
    ap.add_argument("--min-qty", type=float)
    ap.add_argument("--source")
    ap.add_argument("--prices", action="store_true")
    ap.add_argument("--currency", default="GBP", choices=["GBP", "EUR"])
    ap.add_argument("--title")
    ap.add_argument("--limit", type=int)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--dsn", default="host=/var/run/postgresql dbname=denchclaw")
    args = ap.parse_args(argv)
    if not args.out and not args.dry_run:
        ap.error("--out is required unless --dry-run")
    for t in csv_arg(args.type):
        if t.capitalize() not in TYPES:
            ap.error(f"--type {t}: use pack, module, cell or container")
    for s in csv_arg(args.source):
        if s.lower() not in SOURCES:
            ap.error(f"--source {s}: use marketplace, dismantler or bulk")

    stock, bulk = load_rows(args.dsn)
    res = build(stock, bulk, args)
    generated = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    total_kwh = res["total"][res["cols"].index("Total kWh")]
    notes = notes_lines(args, len(res["rows"]), total_kwh, generated)
    cells = [c for r in res["rows"] for c in r] + [c for line in notes for c in line]
    offenders = leak_check(cells, res["private"])
    summary = {
        "file": None if (args.dry_run or offenders) else args.out,
        "rows": len(res["rows"]),
        "total_kwh": round(total_kwh, 1),
        "by_type": dict(collections.Counter(it["Type"] for it in res["items"])),
        "by_country": dict(collections.Counter(it["Country"] or "(blank)" for it in res["items"]).most_common()),
        "excluded": dict(res["excluded"].most_common()),
        "gaps": res["gaps"],
        "leak_check": "ok" if not offenders else "failed",
    }
    if offenders:
        summary["leak_offenders"] = offenders[:200]
        print(json.dumps(summary, indent=2, ensure_ascii=False))
        return 2
    if not args.dry_run:
        write_workbook(args.out, res["cols"], res["rows"], res["total"], notes, args.title)
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
