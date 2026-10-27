"""Tests for stock_list.py. No network, no database: rows are fixtures and load_rows is patched."""
import argparse
import copy
import json

import pytest

import stock_list as sl


def args(**kw):
    base = dict(chemistry=None, type=None, make=None, min_kwh=None, max_kwh=None, countries=None,
                min_qty=None, source=None, prices=False, currency="GBP", title=None, limit=None)
    base.update(kw)
    return argparse.Namespace(**base)


MARKETPLACE = {
    "id": "m1", "status": "listed", "marketplace_status": "published", "marketplace_listing_id": "L1",
    "source": "marketplace", "title": "Audi Q4 82kWh", "quantity": 1, "quantity_unit": "pack",
    "price": 4200, "currency": "GBP", "location": "Waregem, Belgium",
    "supplier_ref": "marketplace:fad2ec32-47c3", "company_name": "Autohandel Didier", "person_name": None,
    "company_country": "Belgium", "company_city": None,
    "details": {"seo_slug": "audi-q4-82kwh-nmc-0b90", "vin": None,
                "marketplace": {"soh": 91, "model": "Q4", "format": "Pack", "pack_kwh": 82.0, "chemistry": "NMC",
                                "manufacturer": "Audi", "year_manufacture": 2022}},
}
DISMANTLER = {
    "id": "d1", "status": "opportunity", "marketplace_status": None, "marketplace_listing_id": None,
    "source": "website", "title": "Tesla Model 3 battery pack (1104422-00-H)", "quantity": 2,
    "quantity_unit": None, "price": 5000, "currency": "EUR",
    "location": "Narėpai (Kaunas district), Lithuania", "supplier_ref": "dalys-31046474",
    "company_name": "GTV Group", "person_name": "Gintas Petraitis", "company_country": "Lithuania",
    "company_city": None,
    "details": {"make": "TESLA", "model": "Model 3", "capacity_kwh": 60, "chemistry": "LiFePO4",
                "condition": "used; tested ok", "tested": True, "year": "2022", "vin": "LRW3E7FA5MC123456",
                "stock_ref": "30114364", "fx": {"date": "2026-10-23", "rate": 1.0, "source": "ECB reference"},
                "ebay": {"seller": "parts_roberto_ltd", "location": {"city": "Kaunas", "country": "LT"}}},
}
HYBRID = {
    "id": "h1", "status": "opportunity", "marketplace_status": None, "marketplace_listing_id": None,
    "source": "supplier list", "title": "Toyota Prius 0.9kWh pack", "quantity": 1, "quantity_unit": "pack",
    "price": None, "currency": None, "location": "Shelford, United Kingdom", "supplier_ref": "12345678",
    "company_name": "LKQ Synetiq", "person_name": None, "company_country": "United Kingdom", "company_city": None,
    "details": {"make": "Toyota", "model": "Prius", "capacity_kwh": 0.9, "chemistry": "NiMH"},
}
STALE = dict(copy.deepcopy(HYBRID), id="s1", status="stale")
MODULE = {
    "id": "x1", "status": "opportunity", "marketplace_status": None, "marketplace_listing_id": None,
    "source": "website", "title": "BMW i3 4.1kWh Module", "quantity": 4, "quantity_unit": None,
    "price": None, "currency": None, "location": "Meppen, Germany", "supplier_ref": "O401894",
    "company_name": "Kempers Autodemontage", "person_name": None, "company_country": "Germany",
    "company_city": None,
    "details": {"make": "BMW", "model": "i3", "capacity_kwh": 4.1, "capacity_ah": 94, "voltage_nominal": 43.75},
}
BULK = {
    "id": "bt_rept", "kind": "supply", "party": "Revive Batteries SL", "contact": "Raul Ortega",
    "summary": "REPT 314Ah prismatic cells, new", "quantity": "7500", "location": "Nevada, United States, US",
    "status": "qualifying", "hold_until": None, "company_name": "Revive Batteries SL", "person_name": None,
    "company_country": "Spain", "company_city": None,
    "spec": {"model": "314Ah", "format": "Cell", "packKwh": 1.0048, "quantity": 7500, "chemistry": "LFP",
             "condition": "new", "manufacturer": "REPT", "value_eur": 226080},
    "context": {},
}
BESS = {
    "id": "bt_bess", "kind": "supply", "party": "Lithos Advisory", "contact": None,
    "summary": "5 unused EP Energy EPES233 liquid-cooled LFP C&I BESS cabinets (233 kWh each)",
    "quantity": "5 units", "location": "Wondelgem (Ghent), Belgium", "status": "open", "hold_until": None,
    "company_name": None, "person_name": None, "company_country": None, "company_city": None,
    "spec": {"trade_kind": "systems"},
    "context": {"listing_title": "5x EP Energy EPES233 233kWh LFP BESS (Damaged)"},
}
CLOSED = dict(copy.deepcopy(BESS), id="bt_closed", status="closed")


def all_stock():
    return [copy.deepcopy(r) for r in (MARKETPLACE, DISMANTLER, HYBRID, STALE, MODULE)]


def all_bulk():
    return [copy.deepcopy(r) for r in (BULK, BESS, CLOSED)]


# --- normalisation -----------------------------------------------------------------------------------

@pytest.mark.parametrize("raw,want", [
    ("LiFePO4", "LFP"), ("LFP", "LFP"), ("lithium iron phosphate", "LFP"), ("NMC622", "NMC 622"),
    ("NMC 622 / graphite", "NMC 622"), ("NCM", "NMC"), ("NMC | UNKNOWN", "NMC"),
    ("LFP | NCA | NMC", ""), ("LMO | NMC", ""), ("GENERIC_LI_ION | UNKNOWN", "Li-ion"),
    ("Li-ion polymer", "Li-ion"), ("NIMH", "NiMH"), ("UNKNOWN", ""), (None, ""),
    ("Manganese-based cathode (LMO family; exact formulation unresolved)", "LMO"),
])
def test_norm_chemistry(raw, want):
    assert sl.norm_chemistry(raw) == want


def test_chemistry_filter_aliases():
    assert sl.chem_filter_value("LiFePO4") == "LFP"
    assert sl.chem_filter_value("ncm") == "NMC"
    assert sl.chem_family("NMC 622") == "NMC"


@pytest.mark.parametrize("raw,want", [
    ("Narėpai (Kaunas district), Lithuania", "Lithuania"), ("Dallas, USA", "United States"),
    ("Belfast, UK (BT3 9DT)", "United Kingdom"), ("Nevada, United States, US", "United States"),
    ("England and Wales", "United Kingdom"), ("Bil1Din North branch", ""), ("Istanbul", ""), (None, ""),
])
def test_norm_country(raw, want):
    assert sl.norm_country(raw) == want


def test_make_normalisation_and_aliases():
    assert sl.norm_make("TESLA") == "Tesla"
    assert sl.norm_make("VW") == "Volkswagen"
    assert sl.make_key("vw") == sl.make_key("Volkswagen")


def test_clean_model_strips_make_year_kwh_and_brackets():
    assert sl.clean_model("Kia", "Kia Sportage NQ5 battery pack (37501-CJ010)") == "Sportage NQ5"
    assert sl.clean_model("Volkswagen", "VW GOLF, E-GOLF VII 2013-20") == "Golf, E-Golf VII"
    assert sl.clean_model("Ford", "FORD KUGA/KUGA PHEV/FHEV 20->>") == "Kuga/Kuga PHEV/FHEV"


def test_condition_vocabulary_hides_free_text():
    assert sl.stock_condition({"condition": "used; Bildelsbasen quality A (TESTAD)"}) == "Used, tested"
    assert sl.stock_condition({"condition": "eBay: For parts or not working"}) == "Damaged / for parts"
    assert sl.stock_condition({"condition": "excellent"}) == "Used, supplier grade Excellent"
    assert sl.stock_condition({"condition": "W"}) == ""
    assert sl.stock_condition({"condition": "used", "installed": True}).endswith("still fitted to vehicle")


# --- type inference ----------------------------------------------------------------------------------

def test_stock_type_inference():
    assert sl.stock_type(MARKETPLACE) == "Pack"
    assert sl.stock_type(MODULE) == "Module"
    unclear = {"title": "Volvo XC40 2023 battery (type unclear)", "details": {"format": "other"}}
    assert sl.stock_type(unclear) == ""
    assert sl.stock_type({"title": "Volvo S60 2020 battery pack", "details": {"format": "other"}}) == "Pack"


def test_bulk_type_inference():
    assert sl.bulk_type(BULK) == "Cell"
    assert sl.bulk_type(BESS) == "Container"
    mods = {"summary": "", "spec": {"trade_kind": "cells"},
            "context": {"listing_title": "120x LG EM048128P6BA 6.6kWh NMC Modules"}}
    assert sl.bulk_type(mods) == "Module"
    mega = {"summary": "", "spec": {"trade_kind": "packs"}, "context": {"listing_title": "1x Tesla Megapack 2.5MWh"}}
    assert sl.bulk_type(mega) == "Container"


def test_bulk_item_parses_title():
    item, reason, priv = sl.build_bulk_item(copy.deepcopy(BESS))
    assert reason is None
    assert (item["Type"], item["Make"], item["Chemistry"], item["kWh per unit"], item["Qty"]) == \
        ("Container", "EP Energy", "LFP", 233.0, 5)
    assert item["Condition / test status"] == "Damaged / for parts"
    assert item["Country"] == "Belgium"
    assert "wondelgem" in priv["towns"] and "lithos advisory" in priv["names"]


# --- filters -----------------------------------------------------------------------------------------

def run_build(**kw):
    return sl.build(all_stock(), all_bulk(), args(**kw))


def test_default_build_excludes_stale_closed_and_tiny_packs():
    res = run_build()
    assert res["excluded"]["stale"] == 1
    assert res["excluded"]["bulk_closed"] == 1
    assert res["excluded"]["pack_under_1kwh"] == 1
    assert {it["_id"] for it in res["items"]} == {"m1", "d1", "x1", "bulk:bt_rept", "bulk:bt_bess"}


def test_min_kwh_zero_keeps_tiny_packs():
    res = run_build(min_kwh=0)
    assert "h1" in {it["_id"] for it in res["items"]}


def test_chemistry_type_make_country_source_filters():
    assert {it["_id"] for it in run_build(chemistry="LiFePO4")["items"]} == {"d1", "bulk:bt_rept", "bulk:bt_bess"}
    assert {it["_id"] for it in run_build(chemistry="LFP", type="cell")["items"]} == {"bulk:bt_rept"}
    assert {it["_id"] for it in run_build(make="tesla,audi")["items"]} == {"m1", "d1"}
    assert {it["_id"] for it in run_build(countries="Belgium")["items"]} == {"m1", "bulk:bt_bess"}
    assert {it["_id"] for it in run_build(source="marketplace")["items"]} == {"m1"}
    assert {it["_id"] for it in run_build(min_kwh=50, max_kwh=70)["items"]} == {"d1"}
    assert {it["_id"] for it in run_build(min_qty=10)["items"]} == {"bulk:bt_rept"}


def test_capacity_class_and_sort_order():
    res = run_build()
    types = [it["Type"] for it in res["items"]]
    assert types == sorted(types, key=sl.TYPES.index)
    audi = next(it for it in res["items"] if it["_id"] == "m1")
    assert audi["kWh per unit"] == 82.0 and audi["Model / product"] == "Q4 e-tron"


def test_refs_are_ours_and_stable():
    a = {it["_id"]: it["Ref"] for it in run_build()["items"]}
    b = {it["_id"]: it["Ref"] for it in run_build()["items"]}
    assert a == b
    assert a["m1"] == "RB-0B90"
    assert all(r.startswith("RB-") for r in a.values())
    assert "31046474" not in " ".join(a.values())


def test_totals_row():
    res = run_build()
    cols = res["cols"]
    total = res["total"]
    assert total[0] == "Total"
    # Qty: 1 + 2 + 4 + 7500 + 5; Total kWh: 82 + 2*60 + 4*4.1 + 7500*1.0048 + 5*233
    assert total[cols.index("Qty")] == 7512
    assert total[cols.index("Total kWh")] == pytest.approx(82 + 120 + 16.4 + 7536 + 1165)


def test_prices_off_by_default_and_marketplace_only():
    res = run_build()
    assert not any(c.startswith("Price") for c in res["cols"])
    res = run_build(prices=True)
    pcol = res["cols"][-1]
    assert pcol == "Price per unit (GBP)"
    prices = {it["_id"]: it[pcol] for it in res["items"]}
    assert prices["m1"] == 4200
    assert prices["d1"] is None                       # dismantler rows never priced
    gap = next(g for g in res["gaps"] if g["column"] == pcol)
    assert gap["by_source"]["dismantler"] == 2 and gap["by_source"]["bulk"] == 2


def test_price_currency_needs_fx():
    item = {"_source": "marketplace", "_price": 4200.0, "_currency": "GBP", "_fx": None}
    assert sl.price_for(item, "GBP") == 4200.0
    assert sl.price_for(item, "EUR") is None
    item["_fx"] = {"rate": 0.84}
    assert sl.price_for(item, "EUR") == 5000.0


def test_column_set_exact():
    assert sl.columns_for(args()) == [
        "Ref", "Type", "Make", "Model / product", "Chemistry", "Format", "kWh per unit", "Ah", "V", "Qty",
        "Total kWh", "SoH %", "Condition / test status", "Year", "Country", "Availability"]
    assert sl.columns_for(args(prices=True, currency="EUR"))[-1] == "Price per unit (EUR)"


def test_gaps_counted_per_column():
    res = run_build()
    gaps = {g["column"]: g["blank"] for g in res["gaps"]}
    assert gaps["SoH %"] == 4                          # only the Audi listing states SoH
    assert "Country" not in gaps


# --- leak check --------------------------------------------------------------------------------------

def private():
    return {"names": sl.name_terms("Parts Roberto Ltd") | sl.name_terms("GTV Group"),
            "towns": sl.town_terms("Narėpai (Kaunas district), Lithuania"),
            "refs": sl.ref_terms("dalys-31046474")}


@pytest.mark.parametrize("cell,why", [
    ("Model 3 from Roberto", "supplier name"), ("Kaunas stock", "town"), ("narėpai", "town"),
    ("see https://dalys.lt/x", "url"), ("www.example.com", "url"), ("sales@example.com", "email"),
    ("+44 7700 900123", "phone"), ("LRW3E7FA5MC123456", "VIN"), ("ref 31046474", "supplier ref"),
    ("EUR 70/kWh", "per-kWh"), ("Row Ash SO32 2HL", "postcode"),
])
def test_leak_check_catches(cell, why):
    offenders = sl.leak_check(["Pack", cell], private())
    assert offenders and why.split()[0].lower() in " ".join(offenders).lower()


def test_leak_check_passes_clean_cells():
    cells = ["RB-0B90", "Pack", "Audi", "Q4 e-tron", "NMC", 82.0, 1, "Lithuania", "Now", 2022,
             "Used, tested", sl.NOTE_TEXT, "Q8 4MN", "CS40HN BESS"]
    assert sl.leak_check(cells, private()) == []


def test_generic_words_are_not_names():
    assert "group" not in sl.name_terms("GTV Group")
    assert "gtv group" in sl.name_terms("GTV Group")


# --- end to end with a patched loader ----------------------------------------------------------------

def test_main_writes_workbook(tmp_path, monkeypatch, capsys):
    from openpyxl import load_workbook
    monkeypatch.setattr(sl, "load_rows", lambda dsn: (all_stock(), all_bulk()))
    out = tmp_path / "s.xlsx"
    assert sl.main(["--out", str(out), "--title", "Test stock"]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["leak_check"] == "ok" and summary["rows"] == 5 and summary["file"] == str(out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Stock", "Notes"]
    ws = wb["Stock"]
    header = [c.value for c in ws[1]]
    assert header == sl.STOCK_COLUMNS
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:P6"
    assert ws[1][0].font.bold
    last = [c.value for c in ws[ws.max_row]]
    assert last[0] == "Total" and last[header.index("Qty")] == 7512
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None).lower()
    for secret in ("didier", "gtv", "roberto", "kaunas", "narėpai", "waregem", "31046474", "lrw3e7fa5mc123456",
                   "30114364", "http", "revive", "nevada", "wondelgem", "ghent"):
        assert secret not in text
    notes = " ".join(str(c.value) for row in wb["Notes"].iter_rows() for c in row if c.value is not None)
    assert sl.NOTE_TEXT in notes and "Generated" in notes


def test_main_dry_run_writes_nothing(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(sl, "load_rows", lambda dsn: (all_stock(), all_bulk()))
    out = tmp_path / "s.xlsx"
    assert sl.main(["--out", str(out), "--dry-run", "--chemistry", "LFP"]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["file"] is None and summary["rows"] == 3 and not out.exists()
    assert summary["by_type"] == {"Pack": 1, "Cell": 1, "Container": 1}


def test_main_leak_fails_hard(tmp_path, monkeypatch, capsys):
    stock = all_stock()
    stock[4]["details"]["model"] = "i3 Meppen stock"                    # town in a model field
    stock[1]["details"]["model"] = "Model 3 via GTV"                    # supplier name
    monkeypatch.setattr(sl, "load_rows", lambda dsn: (stock, all_bulk()))
    out = tmp_path / "s.xlsx"
    assert sl.main(["--out", str(out)]) == 2
    summary = json.loads(capsys.readouterr().out)
    assert summary["leak_check"] == "failed" and not out.exists()
    joined = " ".join(summary["leak_offenders"]).lower()
    assert "meppen" in joined and "gtv" in joined


def test_make_withheld_when_manufacturer_is_the_supplier():
    row = copy.deepcopy(MODULE)
    row["details"]["make"] = "PYLONTECH"
    row["company_name"] = "Pylontech"
    item, reason, priv = sl.build_stock_item(row)
    assert reason is None and item["Make"] == "" and item["_make_withheld"]
    item, _, _ = sl.build_stock_item(copy.deepcopy(MODULE))
    assert item["Make"] == "BMW" and not item["_make_withheld"]


def test_mild_hybrid_excluded_by_wording():
    row = copy.deepcopy(DISMANTLER)
    row["details"]["model"] = "Tiguan III 1.5 eTSI"
    assert sl.build_stock_item(row)[1] == "mild_hybrid_48v"
